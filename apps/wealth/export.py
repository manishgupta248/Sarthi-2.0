"""Excel export of all accounts in one sheet."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Account
from .services import ZERO, latest_snapshot_map

# Indian lakh/crore grouping for Excel
INR_FORMAT = r'[>=10000000]##\,##\,##\,##0.00;[>=100000]##\,##\,##0.00;##,##0.00'
HEADERS = [
    "Holder", "Joint holders", "Account type", "Institution", "Nickname",
    "Account number", "IFSC", "Customer ID", "Branch", "Opening date",
    "Status", "Closed on", "Nominee", "Latest balance (INR)", "Balance as of",
    "Documents", "Remarks",
]


def _safe(text):
    """Stops spreadsheet 'formula injection': text that starts like a formula is neutralised."""
    text = "" if text is None else str(text)
    if text[:1] in ("=", "+", "-", "@"):
        return "'" + text
    return text


def build_accounts_workbook(include_sensitive):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Accounts"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    latest = latest_snapshot_map()
    accounts = Account.objects.select_related(
        "holder", "account_type", "institution", "nominee"
    ).prefetch_related("joint_holders", "documents", "remark_entries")
    total = ZERO
    for account in accounts:
        snapshot = latest.get(account.pk)
        amount = snapshot.amount if snapshot else None
        if amount is not None and not account.is_closed:
            total += amount
        remark_texts = [entry.text for entry in account.remark_entries.all()]
        if include_sensitive:
            number, customer = account.account_number, account.customer_id
            remarks = " | ".join(remark_texts)
        else:
            number = ("XXXXXX" + account.account_number_last4) if account.account_number_last4 else ""
            customer = "(hidden)" if account.customer_id else ""
            remarks = "(hidden)" if remark_texts else ""
        sheet.append([
            _safe(account.holder.full_name),
            _safe(", ".join(person.full_name for person in account.joint_holders.all())),
            _safe(account.account_type.name),
            _safe(account.institution.name),
            _safe(account.nickname),
            _safe(number),
            _safe(account.ifsc_code),
            _safe(customer),
            _safe(account.branch),
            account.opening_date,
            "Closed" if account.is_closed else "Open",
            account.closed_on,
            _safe(account.nominee.full_name if account.nominee else ""),
            float(amount) if amount is not None else None,
            snapshot.as_of if snapshot else None,
            account.documents.count(),
            _safe(remarks),
        ])

    sheet.append([])
    sheet.append(["TOTAL (open accounts)"] + [None] * 12 + [float(total)])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)

    for row in sheet.iter_rows(min_row=2):
        for index in (9, 11, 14):  # opening date, closed on, balance as of
            row[index].number_format = "DD-MM-YYYY"
        row[13].number_format = INR_FORMAT
    sheet.freeze_panes = "A2"
    for index, header in enumerate(HEADERS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(14, len(header) + 2)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
