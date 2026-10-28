"""Excel export of all investments and liabilities in one sheet."""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .export import INR_FORMAT, _safe
from .models import Holding
from .networth import ZERO, holding_value, is_active_on, valuation_map

HEADERS = [
    "Type", "Asset or liability", "Holder", "Joint holder", "Nominee", "Institution or lender",
    "Name", "Number", "Amount (principal, cost or limit)", "Rate % per year", "Start date",
    "End or maturity date", "Next due date", "Instalment, EMI or premium", "Sum assured",
    "Status", "Current value or amount owed", "Value basis", "Remarks",
]
AMOUNT_COLUMNS = ("Amount (principal, cost or limit)", "Instalment, EMI or premium",
                  "Sum assured", "Current value or amount owed")
DATE_COLUMNS = ("Start date", "End or maturity date", "Next due date")
BASIS_TEXT = {"recorded": "Recorded", "estimated": "Estimated", "cost": "Original amount", "none": "None"}


def build_holdings_workbook(include_sensitive):
    today = date.today()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Investments and liabilities"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    valuations = valuation_map()
    assets_total = liabilities_total = ZERO
    holdings = Holding.objects.select_related("holder", "joint_holder", "nominee", "institution")
    for holding in holdings:
        value, source = None, ""
        if is_active_on(holding, today):
            result = holding_value(holding, valuations.get(holding.pk, []), today)
            value, source = result.amount, BASIS_TEXT[result.source]
            if holding.is_liability:
                liabilities_total += value
            else:
                assets_total += value
        if include_sensitive:
            number, remarks = holding.identifier, holding.remarks
        else:
            number = ("XXXXXX" + holding.identifier_last4) if holding.identifier_last4 else ""
            remarks = "(hidden)" if holding.remarks else ""
        sheet.append([
            _safe(holding.get_kind_display()),
            "Liability" if holding.is_liability else "Asset",
            _safe(holding.holder.full_name),
            _safe(holding.joint_holder.full_name if holding.joint_holder else ""),
            _safe(holding.nominee.full_name if holding.nominee else ""),
            _safe(holding.institution.name if holding.institution else holding.counterparty),
            _safe(holding.nickname),
            _safe(number),
            float(holding.principal) if holding.principal is not None else None,
            float(holding.interest_rate) if holding.interest_rate is not None else None,
            holding.start_date,
            holding.end_date,
            holding.next_due_date,
            float(holding.installment_amount) if holding.installment_amount is not None else None,
            float(holding.sum_assured) if holding.sum_assured is not None else None,
            holding.get_status_display(),
            float(value) if value is not None else None,
            source,
            _safe(remarks),
        ])

    value_column = HEADERS.index("Current value or amount owed")
    for label, amount in (("TOTAL ASSETS (investments, not bank accounts)", assets_total),
                          ("TOTAL LIABILITIES", liabilities_total)):
        row = [None] * len(HEADERS)
        row[0] = label
        row[value_column] = float(amount)
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)

    for row in sheet.iter_rows(min_row=2):
        for name in AMOUNT_COLUMNS:
            row[HEADERS.index(name)].number_format = INR_FORMAT
        for name in DATE_COLUMNS:
            row[HEADERS.index(name)].number_format = "DD-MM-YYYY"
    sheet.freeze_panes = "A2"
    for index, header in enumerate(HEADERS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(14, min(len(header) + 2, 30))

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
