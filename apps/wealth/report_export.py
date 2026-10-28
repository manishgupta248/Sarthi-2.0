"""Excel report for one year: monthly summary plus category-by-month sheets."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .export import INR_FORMAT
from .reports import category_matrix, monthly_summary, year_label

HEADER_FILL = PatternFill("solid", fgColor="305496")


def _style_header(sheet):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL


def _number_columns(sheet, first_column):
    for row in sheet.iter_rows(min_row=2):
        for cell in row[first_column:]:
            if isinstance(cell.value, float):
                cell.number_format = INR_FORMAT


def _widths(sheet, count, first=24, other=14):
    for index in range(1, count + 1):
        sheet.column_dimensions[get_column_letter(index)].width = first if index == 1 else other


def build_report_workbook(year, basis, person_id=None, account_id=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Monthly summary"
    sheet.append(["Month", "Income", "Expenses", "Invested", "Left over",
                  "Savings rate %", "Not classified (net out)", "Transactions"])
    rows, totals = monthly_summary(year, basis, person_id, account_id)
    for row in rows + [dict(totals, label="TOTAL " + year_label(year, basis))]:
        sheet.append([
            row["label"], float(row["income"]), float(row["expenses"]), float(row["invested"]),
            float(row["leftover"]), row["savings_rate"], float(row["unclassified"]), row["count"],
        ])
    _style_header(sheet)
    _number_columns(sheet, 1)
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    _widths(sheet, 8)

    for group, title in (
        ("income", "Income by category"),
        ("expense", "Expenses by category"),
        ("investment", "Investments by category"),
    ):
        matrix = category_matrix(group, year, basis, person_id, account_id)
        page = workbook.create_sheet(title)
        page.append(["Category"] + matrix["columns"] + ["Total"])
        for row in matrix["rows"]:
            page.append([row["name"]] + [float(value) for value in row["values"]] + [float(row["total"])])
        page.append(["TOTAL"] + [float(value) for value in matrix["totals"]] + [float(matrix["grand_total"])])
        _style_header(page)
        _number_columns(page, 1)
        for cell in page[page.max_row]:
            cell.font = Font(bold=True)
        _widths(page, len(matrix["columns"]) + 2)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
