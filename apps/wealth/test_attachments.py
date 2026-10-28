import io
import tempfile
from pathlib import Path

from cryptography.fernet import Fernet
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import TestCase, override_settings
from openpyxl import load_workbook

from apps.family.models import Person
from apps.vault.models import Document
from apps.wealth.models import Account, AccountRemark, AccountType, Institution


class AccountUpdateTests(TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.addCleanup(tmp.cleanup)
        override = override_settings(
            VAULT_DIR=Path(tmp.name) / "vault",
            FIELD_ENCRYPTION_KEY=Fernet.generate_key().decode(),
        )
        override.enable()
        self.addCleanup(override.disable)
        User = get_user_model()
        User.objects.create_user(username="boss", password="Test-pass-12345", is_staff=True)
        self.client.login(username="boss", password="Test-pass-12345")
        self.person = Person.objects.create(full_name="Asha Holder")
        self.other = Person.objects.create(full_name="Ravi Joint")
        self.kind = AccountType.objects.get_or_create(name="Savings")[0]
        self.bank = Institution.objects.create(name="Test Bank")
        self.account = Account.objects.create(
            holder=self.person, account_type=self.kind, institution=self.bank, nickname="Main",
        )

    def _form(self, **extra):
        data = {
            "holder": self.person.pk, "account_type": self.kind.pk,
            "institution": self.bank.pk, "nickname": "Main",
        }
        data.update(extra)
        return data

    def _pdf(self, name="passbook.pdf"):
        return SimpleUploadedFile(name, b"%PDF-1.7 passbook data")

    # ----- joint holder -----
    def test_joint_holder_can_be_set_and_removed(self):
        url = f"/wealth/accounts/{self.account.pk}/edit/"
        self.client.post(url, self._form(joint_holder=self.other.pk))
        self.assertEqual(list(self.account.joint_holders.all()), [self.other])
        self.client.post(url, self._form(joint_holder=""))
        self.assertEqual(self.account.joint_holders.count(), 0)

    def test_main_holder_cannot_also_be_the_joint_holder(self):
        url = f"/wealth/accounts/{self.account.pk}/edit/"
        response = self.client.post(url, self._form(joint_holder=self.person.pk))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.account.joint_holders.count(), 0)

    # ----- remarks -----
    def test_first_remark_is_saved_with_a_new_account(self):
        response = self.client.post("/wealth/accounts/add/", self._form(
            nickname="With remark", first_remark="Opened with first salary",
        ))
        self.assertEqual(response.status_code, 302)
        created = Account.objects.get(nickname="With remark")
        self.assertEqual(created.remark_entries.get().text, "Opened with first salary")

    def test_remarks_can_be_added_and_archived(self):
        add = f"/wealth/accounts/{self.account.pk}/remarks/add/"
        self.client.post(add, {"text": "First note"})
        self.client.post(add, {"text": "Second note"})
        self.assertEqual(self.account.remark_entries.count(), 2)
        remark = self.account.remark_entries.first()
        self.client.post(f"/wealth/remarks/{remark.pk}/delete/")
        self.assertEqual(self.account.remark_entries.count(), 1)
        self.assertTrue(AccountRemark.all_objects.filter(pk=remark.pk, is_deleted=True).exists())

    def test_remark_text_is_encrypted_in_the_database(self):
        remark = AccountRemark.objects.create(account=self.account, text="Secret remark ABCDE")
        with connection.cursor() as cursor:
            cursor.execute("SELECT text FROM wealth_accountremark WHERE id = %s", [remark.pk])
            raw = cursor.fetchone()[0]
        self.assertNotIn("Secret", raw)

    def test_account_page_shows_remarks_and_documents(self):
        AccountRemark.objects.create(account=self.account, text="Visible remark")
        self.client.post(f"/wealth/accounts/{self.account.pk}/documents/upload/", {
            "file": self._pdf(), "title": "Passbook page", "doc_type": "bank", "expiry_date": "",
        })
        response = self.client.get(f"/wealth/accounts/{self.account.pk}/")
        self.assertContains(response, "Visible remark")
        self.assertContains(response, "Passbook page")

    # ----- documents -----
    def test_upload_from_the_account_page_stores_and_links_the_document(self):
        response = self.client.post(f"/wealth/accounts/{self.account.pk}/documents/upload/", {
            "file": self._pdf(), "title": "Passbook page", "doc_type": "bank", "expiry_date": "",
        })
        self.assertEqual(response.status_code, 302)
        document = self.account.documents.get()
        self.assertEqual(document.title, "Passbook page")
        self.assertEqual(document.person, self.person)

    def test_several_documents_can_be_added_one_by_one(self):
        url = f"/wealth/accounts/{self.account.pk}/documents/upload/"
        for name in ("first.pdf", "second.pdf", "third.pdf"):
            self.client.post(url, {"file": self._pdf(name), "title": "", "doc_type": "bank", "expiry_date": ""})
        self.assertEqual(self.account.documents.count(), 3)

    def test_bad_file_is_not_stored_or_linked(self):
        self.client.post(f"/wealth/accounts/{self.account.pk}/documents/upload/", {
            "file": SimpleUploadedFile("evil.exe", b"MZ\x90\x00"),
            "title": "", "doc_type": "bank", "expiry_date": "",
        })
        self.assertEqual(Document.objects.count(), 0)
        self.assertEqual(self.account.documents.count(), 0)

    def test_unlink_keeps_the_document_in_the_vault(self):
        self.client.post(f"/wealth/accounts/{self.account.pk}/documents/upload/", {
            "file": self._pdf(), "title": "Passbook", "doc_type": "bank", "expiry_date": "",
        })
        document = self.account.documents.get()
        self.client.post(f"/wealth/accounts/{self.account.pk}/documents/{document.pk}/unlink/")
        self.assertEqual(self.account.documents.count(), 0)
        self.assertEqual(Document.objects.count(), 1)

    # ----- export -----
    def _sheet_values(self, response):
        sheet = load_workbook(io.BytesIO(response.content))["Accounts"]
        return [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None]

    def test_export_hides_remarks_unless_full_details_are_chosen(self):
        AccountRemark.objects.create(account=self.account, text="Private remark")
        masked = self._sheet_values(self.client.post("/wealth/export/", {}))
        self.assertNotIn("Private remark", masked)
        self.assertIn("(hidden)", masked)
        full = self._sheet_values(self.client.post("/wealth/export/", {"full": "1"}))
        self.assertIn("Private remark", full)
