import io
import tempfile
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

from cryptography.fernet import Fernet
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase, override_settings
from openpyxl import load_workbook

from apps.core.models import AuditLog
from apps.family.models import Person
from apps.wealth.models import (
    Account, AccountType, BalanceSnapshot, Holding, HoldingValuation, Institution,
)
from apps.wealth.networth import (
    add_months, estimated_value, holding_attention, holding_value, net_worth_summary,
    net_worth_trend, upcoming_events,
)

KEY = Fernet.generate_key().decode()


def fixed_deposit(**extra):
    values = dict(
        kind="fixed_deposit", principal=Decimal("100000"), interest_rate=Decimal("8"),
        compounding="quarterly", interest_payout="cumulative",
        start_date=date(2026, 1, 1), end_date=date(2027, 1, 1),
    )
    values.update(extra)
    return Holding(**values)


class EstimateTests(SimpleTestCase):
    def test_fixed_deposit_compounds_quarterly(self):
        self.assertEqual(estimated_value(fixed_deposit(), date(2027, 1, 1)), Decimal("108243.22"))

    def test_value_stops_growing_after_maturity(self):
        self.assertEqual(estimated_value(fixed_deposit(), date(2030, 1, 1)), Decimal("108243.22"))

    def test_no_estimate_before_the_start_date(self):
        self.assertIsNone(estimated_value(fixed_deposit(), date(2025, 12, 31)))

    def test_interest_paid_out_keeps_the_principal(self):
        deposit = fixed_deposit(interest_payout="periodic")
        self.assertEqual(estimated_value(deposit, date(2026, 12, 1)), Decimal("100000"))

    def test_simple_interest(self):
        deposit = fixed_deposit(compounding="simple")
        self.assertEqual(estimated_value(deposit, date(2027, 1, 1)), Decimal("108000.00"))

    def _rd(self):
        return Holding(
            kind="recurring_deposit", installment_amount=Decimal("1000"), interest_rate=Decimal("0"),
            start_date=date(2026, 1, 1), end_date=date(2027, 1, 1),
        )

    def test_recurring_deposit_adds_instalments_paid_so_far(self):
        self.assertEqual(estimated_value(self._rd(), date(2026, 3, 15)), Decimal("3000.00"))

    def test_recurring_deposit_is_capped_at_the_term(self):
        self.assertEqual(estimated_value(self._rd(), date(2030, 1, 1)), Decimal("12000.00"))

    def test_recorded_value_beats_the_estimate(self):
        recorded = HoldingValuation(as_of=date(2026, 6, 1), amount=Decimal("105000.00"))
        before = holding_value(fixed_deposit(), [recorded], date(2026, 5, 1))
        after = holding_value(fixed_deposit(), [recorded], date(2026, 7, 1))
        self.assertEqual(before.source, "estimated")
        self.assertEqual((after.amount, after.source), (Decimal("105000.00"), "recorded"))

    def test_fallbacks_for_other_kinds(self):
        today = date(2026, 6, 1)
        bond = Holding(kind="bond", principal=Decimal("5000"))
        loan = Holding(kind="home_loan", principal=Decimal("900000"))
        policy = Holding(kind="insurance", sum_assured=Decimal("1000000"))
        self.assertEqual(holding_value(bond, [], today).source, "cost")
        self.assertEqual(holding_value(loan, [], today).amount, Decimal("900000"))
        self.assertEqual(holding_value(policy, [], today), (Decimal("0.00"), "none", None))

    def test_add_months_handles_short_months(self):
        self.assertEqual(add_months(date(2026, 1, 31), 1), date(2026, 2, 28))
        self.assertEqual(add_months(date(2026, 11, 30), 3), date(2027, 2, 28))


@override_settings(FIELD_ENCRYPTION_KEY=KEY)
class NetWorthTests(TestCase):
    def setUp(self):
        self.asha = Person.objects.create(full_name="Asha Holder")
        kind = AccountType.objects.get_or_create(name="Savings")[0]
        bank = Institution.objects.create(name="Test Bank")
        self.account = Account.objects.create(
            holder=self.asha, account_type=kind, institution=bank, nickname="Main",
        )

    def _holding(self, kind, **extra):
        return Holding.objects.create(kind=kind, holder=self.asha, **extra)

    def _value(self, holding, as_of, amount):
        return HoldingValuation.objects.create(holding=holding, as_of=as_of, amount=Decimal(amount))

    def test_net_worth_adds_assets_and_subtracts_liabilities(self):
        BalanceSnapshot.objects.create(account=self.account, as_of=date(2026, 6, 30), amount=Decimal("10000"))
        deposit = self._holding("fixed_deposit", principal=Decimal("50000"), start_date=date(2026, 1, 1))
        loan = self._holding("home_loan", principal=Decimal("20000"))
        self._value(deposit, date(2026, 6, 1), "52000")
        self._value(loan, date(2026, 6, 1), "15000")
        summary = net_worth_summary(today=date(2026, 7, 1))
        self.assertEqual(summary["assets_total"], Decimal("62000"))
        self.assertEqual(summary["liabilities_total"], Decimal("15000"))
        self.assertEqual(summary["net_worth"], Decimal("47000"))
        self.assertEqual(summary["people"][0]["amount"], Decimal("47000"))

    def test_closed_holdings_are_left_out(self):
        self._holding("bond", principal=Decimal("1000"), status="closed", closed_on=date(2026, 3, 1))
        self._holding("bond", principal=Decimal("500"))
        self.assertEqual(net_worth_summary(today=date(2026, 7, 1))["assets_total"], Decimal("500"))

    def test_trend_starts_when_the_holding_starts(self):
        deposit = self._holding("fixed_deposit", principal=Decimal("100000"), start_date=date(2026, 9, 10))
        self._value(deposit, date(2026, 9, 10), "100000")
        trend = net_worth_trend(months=3, today=date(2026, 10, 27))
        self.assertEqual(trend["values"], [0.0, 100000.0, 100000.0])

    def test_upcoming_dates_and_overdue_items(self):
        today = date(2026, 7, 1)
        self._holding("insurance", nickname="Soon", next_due_date=today + timedelta(days=10))
        self._holding("insurance", nickname="Far", next_due_date=today + timedelta(days=90))
        self._holding("fixed_deposit", nickname="Late", principal=Decimal("1"), end_date=date(2026, 6, 20))
        events = upcoming_events(today)
        names = {event["holding"].nickname: event for event in events}
        self.assertEqual(set(names), {"Soon", "Late"})
        self.assertFalse(names["Soon"]["overdue"])
        self.assertTrue(names["Late"]["overdue"])

    def test_attention_for_missing_and_old_values(self):
        today = date(2026, 12, 1)
        self._holding("gold", nickname="Gold without cost")
        self._holding("insurance", nickname="Term policy")
        old = self._holding("property", nickname="Old value", principal=Decimal("1"))
        self._value(old, date(2026, 1, 1), "5000000")
        summary = net_worth_summary(today=today)
        flagged = {item["holding"].nickname for item in holding_attention(summary["rows"], today)}
        self.assertEqual(flagged, {"Gold without cost", "Old value"})


class PageBase(TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.addCleanup(tmp.cleanup)
        override = override_settings(VAULT_DIR=Path(tmp.name) / "vault", FIELD_ENCRYPTION_KEY=KEY)
        override.enable()
        self.addCleanup(override.disable)
        User = get_user_model()
        User.objects.create_user(username="boss", password="Test-pass-12345", is_staff=True)
        User.objects.create_user(username="normal", password="Test-pass-12345")
        self.client.login(username="boss", password="Test-pass-12345")
        self.asha = Person.objects.create(full_name="Asha Holder")
        self.ravi = Person.objects.create(full_name="Ravi Nominee")
        self.bank = Institution.objects.create(name="Test Bank")


class HoldingPageTests(PageBase):
    def _fd_data(self, **extra):
        data = {
            "holder": self.asha.pk, "institution": self.bank.pk, "nickname": "FD one",
            "new_identifier": "FD-998877", "principal": "100000", "interest_rate": "7.5",
            "compounding": "quarterly", "interest_payout": "cumulative",
            "start_date": "2026-01-01", "end_date": "2027-01-01",
        }
        data.update(extra)
        return data

    def test_add_a_fixed_deposit(self):
        response = self.client.post("/wealth/holdings/add/fixed_deposit/", self._fd_data())
        self.assertEqual(response.status_code, 302)
        deposit = Holding.objects.get()
        self.assertEqual(deposit.kind, "fixed_deposit")
        self.assertEqual(deposit.identifier_last4, "8877")

    def test_required_fields_for_a_deposit(self):
        response = self.client.post("/wealth/holdings/add/fixed_deposit/", self._fd_data(principal=""))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Holding.objects.count(), 0)

    def test_unknown_type_is_not_found(self):
        self.assertEqual(self.client.get("/wealth/holdings/add/nonsense/").status_code, 404)

    def test_add_forms_show_only_relevant_fields(self):
        loan_form = self.client.get("/wealth/holdings/add/loan_given/").content.decode()
        self.assertIn("Borrower", loan_form)
        self.assertNotIn("Sum assured", loan_form)
        policy_form = self.client.get("/wealth/holdings/add/insurance/").content.decode()
        self.assertIn("Premium amount", policy_form)

    def test_loan_given_needs_a_borrower(self):
        response = self.client.post("/wealth/holdings/add/loan_given/", {
            "holder": self.asha.pk, "counterparty": "", "principal": "25000",
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Holding.objects.count(), 0)
        ok = self.client.post("/wealth/holdings/add/loan_given/", {
            "holder": self.asha.pk, "counterparty": "Mr Borrower", "principal": "25000",
        })
        self.assertEqual(ok.status_code, 302)

    def test_joint_holder_cannot_be_the_main_holder(self):
        response = self.client.post(
            "/wealth/holdings/add/fixed_deposit/", self._fd_data(joint_holder=self.asha.pk)
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Holding.objects.count(), 0)

    def test_end_date_cannot_be_before_start_date(self):
        response = self.client.post(
            "/wealth/holdings/add/fixed_deposit/", self._fd_data(end_date="2025-01-01")
        )
        self.assertEqual(response.status_code, 200)

    def test_number_is_masked_and_reveal_is_audited(self):
        self.client.post("/wealth/holdings/add/insurance/", {
            "holder": self.asha.pk, "nominee": self.ravi.pk, "new_identifier": "POL-1234-5678-9012",
            "sum_assured": "1000000", "installment_amount": "12000",
            "installment_frequency": "yearly", "next_due_date": "2027-03-01",
        })
        policy = Holding.objects.get()
        page = self.client.get(f"/wealth/holdings/{policy.pk}/")
        self.assertContains(page, "XXXXXX9012")
        self.assertNotContains(page, "POL-1234-5678-9012")
        edit = self.client.get(f"/wealth/holdings/{policy.pk}/edit/")
        self.assertNotContains(edit, "POL-1234-5678-9012")
        reveal = self.client.post(f"/wealth/holdings/{policy.pk}/reveal/")
        self.assertEqual(reveal.json()["account_number"], "POL-1234-5678-9012")
        self.assertTrue(AuditLog.objects.filter(action="reveal", model_label="wealth.Holding").exists())

    def test_recorded_values_one_per_day(self):
        deposit = Holding.objects.create(kind="gold", holder=self.asha, nickname="Gold")
        url = f"/wealth/holdings/{deposit.pk}/values/add/"
        self.client.post(url, {"as_of": "2026-06-01", "amount": "100000", "note": ""})
        self.client.post(url, {"as_of": "2026-06-01", "amount": "200000", "note": ""})
        self.assertEqual(HoldingValuation.objects.filter(holding=deposit).count(), 1)

    def test_update_values_page_saves_amounts(self):
        deposit = Holding.objects.create(kind="gold", holder=self.asha, nickname="Gold")
        response = self.client.post("/wealth/values/update/", {
            "as_of": "2026-06-01", f"amount_{deposit.pk}": "1,50,000",
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(HoldingValuation.objects.get().amount, Decimal("150000.00"))
        bad = self.client.post("/wealth/values/update/", {
            "as_of": "2026-06-02", f"amount_{deposit.pk}": "abc",
        })
        self.assertEqual(bad.status_code, 200)

    def test_marking_a_premium_paid_moves_the_due_date(self):
        policy = Holding.objects.create(
            kind="insurance", holder=self.asha, nickname="Policy",
            installment_frequency="monthly", next_due_date=date(2026, 1, 31),
        )
        self.client.post(f"/wealth/holdings/{policy.pk}/advance/")
        policy.refresh_from_db()
        self.assertEqual(policy.next_due_date, date(2026, 2, 28))

    def test_single_premium_has_no_next_date_to_move(self):
        policy = Holding.objects.create(
            kind="insurance", holder=self.asha, nickname="Single",
            installment_frequency="single", next_due_date=date(2026, 1, 31),
        )
        self.client.post(f"/wealth/holdings/{policy.pk}/advance/")
        policy.refresh_from_db()
        self.assertEqual(policy.next_due_date, date(2026, 1, 31))

    def test_documents_can_be_uploaded_and_unlinked(self):
        deposit = Holding.objects.create(kind="gold", holder=self.asha, nickname="Gold")
        self.client.post(f"/wealth/holdings/{deposit.pk}/documents/upload/", {
            "file": SimpleUploadedFile("receipt.pdf", b"%PDF-1.7 receipt"),
            "title": "Purchase receipt", "doc_type": "investment", "expiry_date": "",
        })
        document = deposit.documents.get()
        self.assertEqual(document.person, self.asha)
        self.client.post(f"/wealth/holdings/{deposit.pk}/documents/{document.pk}/unlink/")
        self.assertEqual(deposit.documents.count(), 0)

    def test_archive_hides_the_record(self):
        deposit = Holding.objects.create(kind="gold", holder=self.asha, nickname="Gold")
        self.client.post(f"/wealth/holdings/{deposit.pk}/delete/")
        self.assertEqual(Holding.objects.count(), 0)
        self.assertTrue(Holding.all_objects.filter(pk=deposit.pk, is_deleted=True).exists())


class HoldingListAndExportTests(PageBase):
    def setUp(self):
        super().setUp()
        self.deposit = Holding.objects.create(
            kind="fixed_deposit", holder=self.asha, nickname="Asha FD", principal=Decimal("1000"),
            identifier="ID-55556666", remarks="Private remark",
            start_date=date(2026, 1, 1), end_date=date(2027, 1, 1), interest_rate=Decimal("7"),
        )
        self.loan = Holding.objects.create(
            kind="car_loan", holder=self.asha, nickname="Car loan", principal=Decimal("500000"),
        )

    def test_investments_and_liabilities_are_listed_separately(self):
        assets = self.client.get("/wealth/investments/")
        self.assertContains(assets, "Asha FD")
        self.assertNotContains(assets, "Car loan")
        debts = self.client.get("/wealth/liabilities/")
        self.assertContains(debts, "Car loan")
        self.assertNotContains(debts, "Asha FD")

    def test_net_worth_page_and_home_card(self):
        page = self.client.get("/wealth/networth/")
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "Net worth")
        self.assertContains(self.client.get("/"), "Open net worth")

    def test_normal_user_is_blocked(self):
        self.client.logout()
        self.client.login(username="normal", password="Test-pass-12345")
        for url in ("/wealth/networth/", "/wealth/investments/", "/wealth/liabilities/",
                    f"/wealth/holdings/{self.deposit.pk}/", "/wealth/values/update/"):
            self.assertEqual(self.client.get(url).status_code, 403)

    def _values(self, response):
        sheet = load_workbook(io.BytesIO(response.content)).active
        return [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None]

    def test_masked_export_hides_numbers_and_remarks(self):
        values = self._values(self.client.post("/wealth/holdings/export/", {}))
        self.assertNotIn("ID-55556666", values)
        self.assertNotIn("Private remark", values)
        self.assertIn("XXXXXX6666", values)
        self.assertTrue(AuditLog.objects.filter(action="export", note__contains="Investments").exists())

    def test_full_export_includes_them(self):
        values = self._values(self.client.post("/wealth/holdings/export/", {"full": "1"}))
        self.assertIn("ID-55556666", values)
        self.assertIn("Private remark", values)
