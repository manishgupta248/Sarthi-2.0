import io
from datetime import date, timedelta
from decimal import Decimal

from cryptography.fernet import Fernet
from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from openpyxl import load_workbook

from apps.core.models import AuditLog
from apps.family.models import Person
from apps.wealth.models import Account, AccountType, Institution, Transaction, TransactionCategory
from apps.wealth.reports import (
    available_years, category_breakdown, category_matrix, headline, monthly_summary,
    top_transactions, with_comparison, base_queryset, year_bounds, yearly_summary,
)

KEY = Fernet.generate_key().decode()


@override_settings(FIELD_ENCRYPTION_KEY=KEY)
class ReportTests(TestCase):
    def setUp(self):
        User = get_user_model()
        User.objects.create_user(username="boss", password="Test-pass-12345", is_staff=True)
        User.objects.create_user(username="normal", password="Test-pass-12345")
        self.client.login(username="boss", password="Test-pass-12345")
        person = Person.objects.create(full_name="Asha Holder")
        kind = AccountType.objects.get_or_create(name="Savings")[0]
        bank = Institution.objects.create(name="Test Bank")
        self.account = Account.objects.create(holder=person, account_type=kind, institution=bank, nickname="Main")
        cat = lambda name: TransactionCategory.objects.get(name=name)
        self.salary, self.household = cat("Salary"), cat("Household")
        self.refund, self.invest = cat("Refund"), cat("Investment")
        self.transfer, self.medical = cat("Transfer"), cat("Medical")

    def _add(self, when, direction, amount, category=None):
        return Transaction.objects.create(
            account=self.account, date=when, direction=direction,
            amount=Decimal(amount), category=category,
        )

    # ----- category groups -----
    def test_default_category_groups(self):
        expected = {"Salary": "income", "Household": "expense", "Refund": "expense",
                    "Investment": "investment", "Transfer": "transfer", "Other": "other"}
        for name, group in expected.items():
            self.assertEqual(TransactionCategory.objects.get(name=name).group, group)

    # ----- headline numbers -----
    def test_headline_counts_each_group_correctly(self):
        day = date(2026, 5, 10)
        self._add(day, "in", "50000", self.salary)
        self._add(day, "out", "20000", self.household)
        self._add(day, "in", "2000", self.refund)       # refund reduces expenses
        self._add(day, "out", "10000", self.invest)
        self._add(day, "out", "5000", self.transfer)    # transfers are ignored
        self._add(day, "in", "5000", self.transfer)
        self._add(day, "out", "700")                     # no category
        result = headline(base_queryset())
        self.assertEqual(result["income"], Decimal("50000"))
        self.assertEqual(result["expenses"], Decimal("18000"))
        self.assertEqual(result["invested"], Decimal("10000"))
        self.assertEqual(result["leftover"], Decimal("22000"))
        self.assertEqual(result["savings_rate"], 44.0)
        self.assertEqual(result["unclassified"], Decimal("700"))
        self.assertEqual(result["uncategorised_count"], 1)

    # ----- financial year -----
    def test_financial_year_runs_april_to_march(self):
        self.assertEqual(year_bounds(2026, "fy"), (date(2026, 4, 1), date(2027, 3, 31)))
        self.assertEqual(year_bounds(2026, "calendar"), (date(2026, 1, 1), date(2026, 12, 31)))

    def test_monthly_summary_places_months_in_the_right_financial_year(self):
        self._add(date(2027, 2, 10), "out", "300", self.household)   # inside FY 2026
        self._add(date(2026, 3, 31), "out", "900", self.household)   # belongs to FY 2025
        rows, totals = monthly_summary(2026, "fy")
        self.assertEqual(len(rows), 12)
        self.assertEqual(rows[0]["label"], "Apr 2026")
        self.assertEqual(rows[-1]["label"], "Mar 2027")
        feb = [row for row in rows if row["label"] == "Feb 2027"][0]
        self.assertEqual(feb["expenses"], Decimal("300"))
        self.assertEqual(totals["expenses"], Decimal("300"))

    def test_yearly_summary_lists_the_years_with_data(self):
        self._add(date(2025, 5, 1), "out", "100", self.household)
        self._add(date(2026, 6, 1), "out", "200", self.household)
        self.assertEqual(available_years("fy"), [2025, 2026])
        rows = yearly_summary("fy")
        self.assertEqual([row["label"] for row in rows], ["FY 2025-26", "FY 2026-27"])
        self.assertEqual(rows[1]["expenses"], Decimal("200"))

    # ----- category reports -----
    def test_breakdown_is_sorted_with_shares(self):
        self._add(date(2026, 5, 1), "out", "300", self.household)
        self._add(date(2026, 5, 2), "out", "100", self.medical)
        items, total = category_breakdown("expense", date(2026, 5, 1), date(2026, 5, 31))
        self.assertEqual([item["name"] for item in items], ["Household", "Medical"])
        self.assertEqual([item["share"] for item in items], [75.0, 25.0])
        self.assertEqual(total, Decimal("400"))

    def test_comparison_with_the_previous_period(self):
        self._add(date(2026, 4, 5), "out", "200", self.household)
        self._add(date(2026, 5, 5), "out", "300", self.household)
        current, _ = category_breakdown("expense", date(2026, 5, 1), date(2026, 5, 31))
        previous, _ = category_breakdown("expense", date(2026, 4, 1), date(2026, 4, 30))
        item = with_comparison(current, previous)[0]
        self.assertEqual(item["change"], Decimal("100"))
        self.assertEqual(item["change_pct"], 50.0)

    def test_category_by_month_table(self):
        self._add(date(2026, 4, 5), "out", "100", self.household)
        self._add(date(2026, 5, 5), "out", "250", self.household)
        matrix = category_matrix("expense", 2026, "fy")
        row = matrix["rows"][0]
        self.assertEqual(row["name"], "Household")
        self.assertEqual(row["values"][:2], [Decimal("100"), Decimal("250")])
        self.assertEqual(matrix["grand_total"], Decimal("350"))

    def test_largest_items_use_the_right_direction(self):
        self._add(date(2026, 5, 1), "out", "50", self.household)
        self._add(date(2026, 5, 2), "out", "900", self.household)
        self._add(date(2026, 5, 3), "in", "9999", self.household)    # refund: not an "expense item"
        top = top_transactions("expense", date(2026, 5, 1), date(2026, 5, 31))
        self.assertEqual([t.amount for t in top], [Decimal("900.00"), Decimal("50.00")])

    # ----- pages -----
    def test_report_pages_load_and_warn_about_uncategorised_entries(self):
        self._add(date.today(), "out", "100")
        monthly = self.client.get("/wealth/reports/")
        self.assertEqual(monthly.status_code, 200)
        self.assertContains(monthly, "no category")
        self.assertEqual(self.client.get("/wealth/reports/?view=yearly").status_code, 200)
        self.assertEqual(self.client.get("/wealth/reports/?basis=calendar").status_code, 200)
        self.assertEqual(self.client.get("/wealth/reports/categories/?group=income").status_code, 200)
        self.assertEqual(self.client.get("/wealth/reports/categories/?group=expense&month=6").status_code, 200)
        self.assertEqual(self.client.get("/wealth/reports/categories/?group=nonsense").status_code, 200)

    def test_excel_report_has_the_expected_sheets_and_is_audited(self):
        self._add(date(2026, 5, 1), "out", "100", self.household)
        response = self.client.post("/wealth/reports/export/", {"year": "2026", "basis": "fy"})
        self.assertEqual(response.status_code, 200)
        names = load_workbook(io.BytesIO(response.content)).sheetnames
        self.assertEqual(names, ["Monthly summary", "Income by category", "Expenses by category", "Investments by category"])
        self.assertTrue(AuditLog.objects.filter(action="export", note__contains="Report export").exists())

    def test_normal_user_cannot_open_reports(self):
        self.client.logout()
        self.client.login(username="normal", password="Test-pass-12345")
        self.assertEqual(self.client.get("/wealth/reports/").status_code, 403)
        self.assertEqual(self.client.post("/wealth/reports/export/", {}).status_code, 403)


@override_settings(FIELD_ENCRYPTION_KEY=KEY)
class BrowsingTests(TestCase):
    def setUp(self):
        User = get_user_model()
        User.objects.create_user(username="boss", password="Test-pass-12345", is_staff=True)
        self.client.login(username="boss", password="Test-pass-12345")
        person = Person.objects.create(full_name="Asha Holder")
        kind = AccountType.objects.get_or_create(name="Savings")[0]
        bank = Institution.objects.create(name="Test Bank")
        self.account = Account.objects.create(holder=person, account_type=kind, institution=bank, nickname="Main")

    def _add(self, when, amount="10"):
        return Transaction.objects.create(
            account=self.account, date=when, direction="out", amount=Decimal(amount)
        )

    def _count(self, url):
        return self.client.get(url).context["page"].paginator.count

    def test_page_opens_on_the_current_month_by_default(self):
        self._add(date.today())
        self._add(date.today() - timedelta(days=400))
        self.assertEqual(self._count("/wealth/transactions/"), 1)
        self.assertEqual(self._count("/wealth/transactions/?period=all"), 2)

    def test_year_and_month_browsing(self):
        self._add(date(2025, 1, 10))
        self._add(date(2025, 2, 10))
        self._add(date(2024, 1, 10))
        self.assertEqual(self._count("/wealth/transactions/?year=2025&month=1"), 1)
        self.assertEqual(self._count("/wealth/transactions/?year=2025"), 2)
        self.assertEqual(self._count("/wealth/transactions/?date_from=2024-01-01&date_to=2024-12-31"), 1)

    def test_uncategorised_filter(self):
        self._add(date(2025, 1, 10))
        self.assertEqual(self._count("/wealth/transactions/?category=none&period=all"), 1)

    def test_archive_overview_lists_years_and_months(self):
        self._add(date(2025, 1, 10))
        self._add(date(2025, 2, 10))
        self._add(date(2024, 6, 10))
        response = self.client.get("/wealth/transactions/archive/")
        years = response.context["years"]
        self.assertEqual([y["year"] for y in years], [2025, 2024])
        self.assertEqual([m["month"] for m in years[0]["months"]], [2, 1])
        self.assertEqual(years[0]["count"], 2)

    def test_category_pages(self):
        self.assertEqual(self.client.get("/wealth/categories/").status_code, 200)
        response = self.client.post("/wealth/categories/add/", {"name": "Pets", "group": "expense", "sort_order": 200})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(TransactionCategory.objects.get(name="Pets").group, "expense")
        duplicate = self.client.post("/wealth/categories/add/", {"name": "pets", "group": "expense", "sort_order": 1})
        self.assertEqual(duplicate.status_code, 200)

    def test_used_category_cannot_be_archived(self):
        category = TransactionCategory.objects.get(name="Household")
        Transaction.objects.create(account=self.account, date=date(2025, 1, 1), direction="out",
                                   amount=Decimal("5"), category=category)
        self.client.post(f"/wealth/categories/{category.pk}/delete/")
        self.assertTrue(TransactionCategory.objects.filter(pk=category.pk).exists())
