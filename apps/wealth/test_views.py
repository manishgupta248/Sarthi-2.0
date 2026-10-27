import io
from datetime import date
from decimal import Decimal

from cryptography.fernet import Fernet
from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from openpyxl import load_workbook

from apps.core.models import AuditLog
from apps.family.models import Person
from apps.wealth.models import Account, AccountType, BalanceSnapshot, Institution

KEY = Fernet.generate_key().decode()
NUMBER = "50100123456789"


@override_settings(FIELD_ENCRYPTION_KEY=KEY)
class WealthViewTests(TestCase):
    def setUp(self):
        User = get_user_model()
        User.objects.create_user(username="boss", password="Test-pass-12345", is_staff=True)
        User.objects.create_user(username="normal", password="Test-pass-12345")
        self.client.login(username="boss", password="Test-pass-12345")
        self.person = Person.objects.create(full_name="Asha Holder")
        self.kind = AccountType.objects.get_or_create(name="Savings")[0]
        self.bank = Institution.objects.create(name="Test Bank")
        self.account = Account.objects.create(
            holder=self.person, account_type=self.kind, institution=self.bank,
            account_number=NUMBER, customer_id="CUST998877", nickname="Main account",
        )

    def _sheet_values(self, response):
        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook["Accounts"]
        return [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None]

    def test_account_page_shows_only_the_masked_number(self):
        response = self.client.get(f"/wealth/accounts/{self.account.pk}/")
        self.assertContains(response, "XXXXXX6789")
        self.assertNotContains(response, NUMBER)
        self.assertNotContains(response, "CUST998877")

    def test_edit_page_never_shows_the_stored_number(self):
        response = self.client.get(f"/wealth/accounts/{self.account.pk}/edit/")
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, NUMBER)

    def test_reveal_is_post_only_staff_only_and_audited(self):
        url = f"/wealth/accounts/{self.account.pk}/reveal/"
        self.assertEqual(self.client.get(url).status_code, 405)
        response = self.client.post(url)
        self.assertEqual(response.json()["account_number"], NUMBER)
        self.assertIn("no-store", response["Cache-Control"])
        self.assertTrue(AuditLog.objects.filter(action="reveal").exists())
        self.client.logout()
        self.client.login(username="normal", password="Test-pass-12345")
        self.assertEqual(self.client.post(url).status_code, 403)

    def test_add_account_through_the_page(self):
        response = self.client.post("/wealth/accounts/add/", {
            "holder": self.person.pk, "account_type": self.kind.pk, "institution": self.bank.pk,
            "nickname": "New one", "new_account_number": "9988 7766 5544",
            "ifsc_code": "sbin0001234", "branch": "", "opening_date": "", "closed_on": "",
            "remarks": "", "new_customer_id": "",
        })
        self.assertEqual(response.status_code, 302)
        created = Account.objects.get(nickname="New one")
        self.assertEqual(created.account_number_last4, "5544")
        self.assertEqual(created.ifsc_code, "SBIN0001234")

    def test_bad_ifsc_is_rejected(self):
        response = self.client.post("/wealth/accounts/add/", {
            "holder": self.person.pk, "account_type": self.kind.pk, "institution": self.bank.pk,
            "nickname": "Bad IFSC", "ifsc_code": "NOT-AN-IFSC",
        })
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Account.objects.filter(nickname="Bad IFSC").exists())

    def test_one_balance_per_day_is_enforced(self):
        url = f"/wealth/accounts/{self.account.pk}/snapshots/add/"
        self.client.post(url, {"as_of": "2026-06-01", "amount": "100.00", "note": ""})
        self.client.post(url, {"as_of": "2026-06-01", "amount": "200.00", "note": ""})
        self.assertEqual(BalanceSnapshot.objects.filter(account=self.account).count(), 1)

    def test_update_balances_page_saves_amounts_with_commas(self):
        response = self.client.post("/wealth/balances/update/", {
            "as_of": "2026-06-01", f"amount_{self.account.pk}": "1,23,456.50",
        })
        self.assertEqual(response.status_code, 302)
        snapshot = BalanceSnapshot.objects.get(account=self.account)
        self.assertEqual(snapshot.amount, Decimal("123456.50"))

    def test_update_balances_rejects_a_bad_amount(self):
        response = self.client.post("/wealth/balances/update/", {
            "as_of": "2026-06-01", f"amount_{self.account.pk}": "abc",
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(BalanceSnapshot.objects.count(), 0)

    def test_masked_export_hides_numbers(self):
        response = self.client.post("/wealth/export/", {})
        self.assertEqual(response.status_code, 200)
        values = self._sheet_values(response)
        self.assertNotIn(NUMBER, values)
        self.assertNotIn("CUST998877", values)
        self.assertIn("XXXXXX6789", values)
        self.assertTrue(AuditLog.objects.filter(action="export").exists())

    def test_full_export_contains_numbers_and_is_audited(self):
        response = self.client.post("/wealth/export/", {"full": "1"})
        self.assertIn(NUMBER, self._sheet_values(response))
        self.assertTrue(AuditLog.objects.filter(action="export", note__contains="FULL").exists())

    def test_export_neutralises_formula_text(self):
        Account.objects.create(
            holder=self.person, account_type=self.kind, institution=self.bank,
            nickname="=HYPERLINK(\"http://evil\")",
        )
        values = self._sheet_values(self.client.post("/wealth/export/", {}))
        self.assertIn("'=HYPERLINK(\"http://evil\")", values)

    def test_normal_user_cannot_open_wealth_pages(self):
        self.client.logout()
        self.client.login(username="normal", password="Test-pass-12345")
        self.assertEqual(self.client.get("/wealth/").status_code, 403)
        self.assertEqual(self.client.get("/wealth/accounts/").status_code, 403)

    def test_wealth_dashboard_and_home_widget_load(self):
        BalanceSnapshot.objects.create(account=self.account, as_of=date.today(), amount=Decimal("5000.00"))
        self.assertContains(self.client.get("/wealth/"), "5,000.00")
        home = self.client.get("/")
        self.assertContains(home, "Open wealth dashboard")
